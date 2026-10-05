"""文档加载：支持 PDF / DOCX / TXT / Markdown，统一产出 langchain Document。

故意不依赖 langchain-community 里的各类 Loader，直接调用 pypdf / python-docx，
依赖更少、行为更可控。
"""

from __future__ import annotations

import logging
from pathlib import Path

from langchain_core.documents import Document

logger = logging.getLogger(__name__)

SUPPORTED_SUFFIXES: set[str] = {".pdf", ".docx", ".txt", ".md", ".markdown"}


def _meta(path: Path, **extra) -> dict:
    return {"source": str(path), "file_name": path.name, "file_type": path.suffix.lower().lstrip("."), **extra}


def load_pdf(path: Path) -> list[Document]:
    """按页加载 PDF，一页一个 Document（便于引用页码）。"""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    docs: list[Document] = []
    for page_no, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            docs.append(Document(page_content=text, metadata=_meta(path, page=page_no)))
    return docs


def load_docx(path: Path) -> list[Document]:
    """加载 Word 文档：正文段落 + 表格内容。"""
    import docx

    document = docx.Document(str(path))
    parts: list[str] = [p.text.strip() for p in document.paragraphs if p.text.strip()]

    # 表格按行拼成 "单元格 | 单元格" 形式，避免信息丢失
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    text = "\n".join(parts).strip()
    return [Document(page_content=text, metadata=_meta(path))] if text else []


def load_text(path: Path) -> list[Document]:
    """加载纯文本 / Markdown，按二级标题切成粗粒度块，保留结构语义。"""
    raw = path.read_text(encoding="utf-8", errors="ignore").strip()
    if not raw:
        return []

    blocks: list[str] = []
    buffer: list[str] = []
    for line in raw.splitlines():
        if line.startswith("## ") and buffer:
            blocks.append("\n".join(buffer).strip())
            buffer = [line]
        else:
            buffer.append(line)
    if buffer:
        blocks.append("\n".join(buffer).strip())
    blocks = [b for b in blocks if b]

    return [Document(page_content=b, metadata=_meta(path, block=i)) for i, b in enumerate(blocks, start=1)]


_LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".txt": load_text,
    ".md": load_text,
    ".markdown": load_text,
}


def load_document(path: str | Path) -> list[Document]:
    """加载单个文件，返回 Document 列表。"""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"文件不存在：{path}")

    loader = _LOADERS.get(path.suffix.lower())
    if loader is None:
        raise ValueError(f"不支持的文件类型 {path.suffix}，当前支持：{', '.join(sorted(SUPPORTED_SUFFIXES))}")

    docs = loader(path)
    logger.info("已加载 %s，共 %d 个片段", path.name, len(docs))
    return docs


def load_directory(directory: str | Path, recursive: bool = True) -> list[Document]:
    """批量加载目录下所有受支持的文件。"""
    directory = Path(directory)
    if not directory.is_dir():
        raise NotADirectoryError(f"目录不存在：{directory}")

    pattern = "**/*" if recursive else "*"
    all_docs: list[Document] = []
    for file in sorted(directory.glob(pattern)):
        if not file.is_file() or file.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        try:
            all_docs.extend(load_document(file))
        except Exception as exc:  # 单个文件失败不影响整体入库
            logger.warning("跳过 %s：%s", file.name, exc)
    return all_docs
