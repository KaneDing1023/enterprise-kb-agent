"""冒烟测试：不消耗 DashScope 额度，仅验证依赖可导入、加载与切分链路正常。

运行：
    python tests/test_smoke.py
或（若已安装 pytest）：
    pytest tests -q
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


def test_dependencies_importable() -> None:
    import chromadb  # noqa: F401
    import dashscope  # noqa: F401
    import langchain_chroma  # noqa: F401
    import streamlit  # noqa: F401
    from langchain_text_splitters import RecursiveCharacterTextSplitter  # noqa: F401


def test_config_loads() -> None:
    from config import settings

    assert settings.embedding_model
    assert settings.llm_model
    assert settings.chunk_size > settings.chunk_overlap
    assert settings.persist_dir.is_absolute()


def test_txt_load_and_split() -> None:
    from src.document_loader import load_document
    from src.text_splitter import split_documents

    with tempfile.TemporaryDirectory() as tmp:
        file = Path(tmp) / "制度.txt"
        file.write_text(
            "## 请假制度\n" + "员工请假需提前一天在系统提交申请。" * 40 +
            "\n## 报销制度\n" + "差旅报销需在行程结束后五个工作日内提交发票。" * 40,
            encoding="utf-8",
        )
        docs = load_document(file)
        assert docs and all(d.metadata["file_name"] == "制度.txt" for d in docs)

        chunks = split_documents(docs, chunk_size=200, chunk_overlap=40)
        assert len(chunks) > len(docs)
        assert all(len(c.page_content) <= 200 for c in chunks)
        assert chunks[0].metadata["chunk_index"] == 0
        assert chunks[-1].metadata["chunk_total"] == len(chunks)


def test_docx_load() -> None:
    import docx

    from src.document_loader import load_document

    with tempfile.TemporaryDirectory() as tmp:
        file = Path(tmp) / "通知.docx"
        document = docx.Document()
        document.add_paragraph("关于中秋节放假的通知")
        document.add_paragraph("9 月 25 日至 27 日放假，共 3 天。")
        document.save(str(file))

        docs = load_document(file)
        assert len(docs) == 1
        assert "中秋节" in docs[0].page_content


def test_unsupported_suffix() -> None:
    from src.document_loader import load_document

    with tempfile.TemporaryDirectory() as tmp:
        file = Path(tmp) / "a.xyz"
        file.write_text("x", encoding="utf-8")
        try:
            load_document(file)
        except ValueError:
            return
        raise AssertionError("应当对不支持的文件类型抛出 ValueError")


if __name__ == "__main__":
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
    failed = 0
    for fn in tests:
        try:
            fn()
            print(f"[PASS] {fn.__name__}")
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print(f"[FAIL] {fn.__name__}: {exc}")
    print(f"\n{len(tests) - failed}/{len(tests)} 通过")
    raise SystemExit(1 if failed else 0)
